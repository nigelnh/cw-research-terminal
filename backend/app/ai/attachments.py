"""Bounded, request-local document extraction. Files are never executed or retained."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from fastapi import HTTPException

MAX_FILES = 2
MAX_FILE_BYTES = 20 * 1024 * 1024  # combined, per request
MAX_EXTRACTED_CHARS = 200_000
MAX_ZIP_BYTES = 40 * 1024 * 1024
SUPPORTED_EXTENSIONS = {'.pdf', '.txt', '.md', '.csv', '.tsv', '.json', '.xlsx', '.docx'}


@dataclass
class Document:
    name: str
    sections: list[tuple[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    characters: int = 0

    def add(self, location: str, text: str) -> bool:
        text = text.strip()
        if not text:
            return True
        if len(self.sections) >= 2000:
            self.warn('Extraction limited to 2,000 text sections; some content was omitted.')
            return False
        remaining = MAX_EXTRACTED_CHARS - self.characters
        if len(text) > remaining:
            self.warn('Extraction limited to 200,000 characters; some content was omitted.')
        text = text[:remaining]
        for offset in range(0, len(text), 1000):
            self.sections.append((location, text[offset:offset + 1000]))
        self.characters += len(text)
        return self.characters < MAX_EXTRACTED_CHARS

    def warn(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)


def safe_name(name: str) -> str:
    return re.sub(r'[\x00-\x1f\x7f]', '', PurePosixPath(name.replace('\\', '/')).name)[:160] or 'document'


def reject(message: str, status: int = 400):
    raise HTTPException(status_code=status, detail=message)


def _checked_archive(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    infos = archive.infolist()
    if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
        archive.close()
        reject('The document expands beyond the supported extraction size.')
    if any(i.flag_bits & 1 for i in infos):
        archive.close()
        reject('Password-protected documents are not supported.')
    return archive


def extract_document(filename: str, data: bytes) -> Document:
    doc = Document(safe_name(filename))
    extension = PurePosixPath(doc.name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        reject('Use PDF, TXT, MD, CSV, TSV, JSON, XLSX or DOCX files.', 415)
    if not data:
        reject(f'{doc.name}: the file is empty.')
    try:
        if extension == '.pdf':
            from pypdf import PdfReader
            if not data.startswith(b'%PDF-'):
                reject(f'{doc.name}: invalid PDF.')
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                reject(f'{doc.name}: password-protected PDFs are not supported.')
            if len(reader.pages) > 100:
                doc.warn('Only the first 100 PDF pages were extracted.')
            for number, page in enumerate(reader.pages[:100], 1):
                # pypdf may need considerable memory for a dense content stream.
                content = page.get_contents()
                if content is not None and len(content.get_data()) > 5 * 1024 * 1024:
                    doc.warn(f'Page {number} was too complex to extract.')
                    continue
                text = page.extract_text() or ''
                if not text.strip():
                    doc.warn(f'Page {number} has no extractable text; scanned content/images were not read.')
                if not doc.add(f'page {number}', text):
                    break
        elif extension == '.docx':
            from defusedxml import ElementTree
            with _checked_archive(data) as archive:
                root = ElementTree.fromstring(archive.read('word/document.xml'))
            ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            for number, para in enumerate(root.iter(f'{ns}p'), 1):
                text = ''.join(node.text or '' for node in para.iter(f'{ns}t'))
                if not doc.add(f'paragraph {number}', text):
                    break
            doc.warn('DOCX text/tables only; embedded images and charts were not read.')
        elif extension == '.xlsx':
            from openpyxl import load_workbook
            with _checked_archive(data):
                pass
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > 12:
                    doc.warn('Only the first 12 worksheets were extracted.')
                for sheet in workbook.worksheets[:12]:
                    if (sheet.max_row or 0) > 3000 or (sheet.max_column or 0) > 40:
                        doc.warn(f'{sheet.title}: extraction limited to 3,000 rows and 40 columns.')
                    header = ''
                    for number, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 3000, 3000), max_col=min(sheet.max_column or 40, 40), values_only=True), 1):
                        values = [str(v) if v is not None else '' for v in row]
                        if not any(values):
                            continue
                        text = ' | '.join(values)
                        if number == 1:
                            header = text[:800]
                        if any(v.startswith('=') for v in values):
                            doc.warn('Spreadsheet formulas are shown as text, not calculated results.')
                        if not doc.add(f'sheet {sheet.title}, row {number}', f'Columns: {header}\n{text}' if number > 1 else text):
                            break
                    if doc.characters >= MAX_EXTRACTED_CHARS:
                        break
            finally:
                workbook.close()
        else:
            encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            text = data.decode(encoding)
            if '\x00' in text:
                reject(f'{doc.name}: binary content is not a supported text document.')
            if extension in ('.csv', '.tsv'):
                reader = csv.reader(io.StringIO(text), delimiter='\t' if extension == '.tsv' else ',')
                header = ''
                for number, row in enumerate(reader, 1):
                    line = ' | '.join(row)
                    if number == 1:
                        header = line[:800]
                    if not doc.add(f'row {number}', f'Columns: {header}\n{line}' if number > 1 else line):
                        break
            else:
                if extension == '.json':
                    text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
                for number, line in enumerate(text.splitlines(), 1):
                    if not doc.add(f'line {number}', line):
                        break
    except HTTPException:
        raise
    except Exception:
        reject(f'{doc.name}: could not read this file. Check its format and remove any password.')
    if not doc.sections:
        reject(f'{doc.name}: no readable text found. For scanned PDFs, upload a text/OCR version.')
    return doc


async def extract_in_worker(filename: str, data: bytes) -> Document:
    import asyncio
    import sys
    process = await asyncio.create_subprocess_exec(
        sys.executable, '-m', 'app.ai.attachment_worker', safe_name(filename),
        stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        output, _ = await asyncio.wait_for(process.communicate(data), timeout=20)
        if process.returncode != 0:
            reject('The document exceeded extraction limits. Try a smaller file.')
        result = json.loads(output)
        if 'error' in result:
            reject(result['error'], result.get('status', 400))
        return Document(**result['document'])
    except TimeoutError:
        reject('Document extraction timed out. Try a smaller file.')
    finally:
        if process.returncode is None:
            process.kill()
        await process.wait()
